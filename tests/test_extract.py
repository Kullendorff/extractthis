"""Regressionstester för extract-this. Kör:  python -m pytest tests -q
(fixtures byggs automatiskt första gången via make_fixtures.py)."""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
FIX = HERE / "fixtures"
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE))

from bench_ocr import cer  # noqa: E402
from xt import common, ocr, pdf  # noqa: E402
from xt.router import extract_file, iter_jsonl, render_md, render_txt  # noqa: E402
from xt.structured import SIGN, Book, to_int, trim_chunk_body  # noqa: E402


@pytest.fixture(scope="session", autouse=True)
def fixtures():
    if not (FIX / "truth.json").exists():
        subprocess.run([sys.executable, str(HERE / "make_fixtures.py")], check=True)
    return json.loads((FIX / "truth.json").read_text(encoding="utf-8"))


def text(doc):
    return "\n".join(s["text"] for s in doc.sections)


# --------------------------------------------------------------------------- common
def test_clean_text_dehyphenation_keeps_och():
    assert common.clean_text("för-\noch nackdelar") == "för- och nackdelar"
    assert common.clean_text("par-\nticularly") == "particularly"
    assert common.clean_text("par­ ticularly") == "particularly"
    assert common.clean_text("ﬁnal") == "final"


def test_known_word_ratio_separates_garbage_from_tables():
    good = "Craft: Corellian YT-1300 Transport Type: Light freighter Scale: Starfighter Length: meters " * 10
    bad = "weren'l around ~auseyouwere friend., lew pwpi.. Inl\"raded IO\"ilier olwdety liIe IUSI " * 10
    assert not common.looks_like_garbage(good)
    assert common.looks_like_garbage(bad)


def test_fantasy_swedish_visible_text_is_not_garbage():
    """EON Mystik & Magi s.14 (0,82 kända ord) OCR:ades felaktigt – synlig text
    har lägre skräpgräns än dolt OCR-lager."""
    t = ("Kabalaorden använder en mängd unika termer: genus (aspekt), intensitet (magnitud), "
         "generering (alstra), formulus (besvärjelse), ritus (ritual), opus (alkemi), magi (magiker), "
         "aqua (hydrotropi), aër (pneumotropi), ignis (pyrotropi), solum (geotropi), Ebhron Thalamur ") * 4
    assert not common.looks_like_garbage(t, visible=True)


def test_layout_dropcap_fix():
    assert pdf._fix_dropcaps("essa regler kan användas. DSom spelledare kan man.") == \
        "Dessa regler kan användas. Som spelledare kan man."
    assert pdf._fix_dropcaps("och sedan NASA Mission") == "och sedan NASA Mission"


def test_parse_pages():
    assert common.parse_pages("1-3,5,9-", 10) == [0, 1, 2, 4, 8, 9]
    assert common.parse_pages(None, 3) == [0, 1, 2]


def test_md_table_escapes_pipes():
    assert "a\\|b" in common.md_table([["x"], ["a|b"]])


def test_junk_filter_keeps_dates_and_personnummer():
    assert not ocr._is_junk("2026-03-14 10:32")
    assert not ocr._is_junk("19121212-1212")
    assert ocr._is_junk("17")
    assert ocr._is_junk("xii")
    assert ocr._is_junk("~~~•••~~~••")


def test_repeated_edges_never_eat_per_page_data():
    """Granskning 2026-10-10: siffer-normalisering raderade datum/blodtryck från
    varje journalsida utom den första."""
    pages = {i: f"Datum: 2024-03-0{i + 1}\nSignerad: Dr A\nUnik mening {i} om patienten.\n"
                f"Mer text här.\nBlodtryck: 12{i}/80\nSida {i + 1}" for i in range(8)}
    out = pdf.strip_repeated_edges(pages)
    for i in range(8):
        assert f"2024-03-0{i + 1}" in out[i] and f"12{i}/80" in out[i]
    assert "Sida 1" in out[0] and "Sida 5" not in out[4]          # riktiga sidnummer bort (utom första)
    assert "Signerad: Dr A" in out[0] and "Signerad: Dr A" not in out[3]


def test_junk_only_at_edges_and_no_false_roman():
    assert not ocr._is_junk("120", at_edge=False)
    assert not ocr._is_junk("mild") and not ocr._is_junk("civil") and not ocr._is_junk("DIM")
    assert not ocr._is_junk("Ja")


def test_noise_is_hidden_not_dropped():
    res = {"words": [{"t": t, "x": 10 * k, "y": 10, "w": 30, "h": 12, "c": 40, "b": 1, "p": 1, "l": 1}
                     for k, t in enumerate("Patienten mår bra idag".split())]}
    md = ocr.ocr_pages_to_markdown({0: res})[0]
    assert "Patienten mår bra idag" in md and md.startswith("<!-- osäker OCR")


def test_dropcap_fix_leaves_acronym_plurals():
    assert pdf._fix_dropcaps("in the old days, PCs were slow") == "in the old days, PCs were slow"
    assert pdf._fix_dropcaps("en ny TVn stod där") == "en ny TVn stod där"


def test_utf16_and_unknown_charset(tmp_path):
    p = tmp_path / "u16.csv"
    p.write_bytes("namn;stad\nÅsa;Göteborg\n".encode("utf-16"))
    assert "| Åsa | Göteborg |" in text(extract_file(p))
    e = tmp_path / "old.eml"
    e.write_bytes(b"Subject: Hej\nContent-Type: text/plain; charset=\"unknown-8bit\"\n\nGammal text \xe5\xe4\xf6\n")
    assert "Gammal text" in text(extract_file(e))


def test_sheet_formatted_empty_rows_not_reported_as_omitted(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(10):
        ws.append([i, f"rad {i}"])
    ws.cell(row=1000, column=1).number_format = "0.00"  # "formaterad" tomrad långt ner
    p = tmp_path / "f.xlsx"
    wb.save(p)
    assert "utelämnade" not in text(extract_file(p))


def test_docx_textbox_and_footnote(tmp_path):
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    d.add_paragraph("Huvudtext")
    W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    box = parse_xml(f'<w:r {W_NS}><w:pict><w:txbxContent><w:p><w:r><w:t>Text i ruta</w:t></w:r></w:p>'
                    f'</w:txbxContent></w:pict></w:r>')
    d.paragraphs[0]._p.append(box)
    p = tmp_path / "tb.docx"
    d.save(p)
    md = text(extract_file(p))
    assert "Huvudtext" in md and "Text i ruta" in md


def test_docm_is_readable(tmp_path):
    import zipfile

    src, dst = FIX / "rapport.docx", tmp_path / "makro.docm"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w") as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml",
                                    b"PLACEHOLDER").replace(b"application/vnd.openxmlformats-officedocument.PLACEHOLDER",
                                                            b"application/vnd.ms-word.document.macroEnabled.main+xml")
            zo.writestr(it, data)
    assert "# Rapport om knäskador" in text(extract_file(dst))


def test_cli_name_collisions_and_combine_self(tmp_path):
    exe = [sys.executable, str(SCRIPTS / "extract.py")]
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir(), b.mkdir()
    (a / "notes.txt").write_text("från a", encoding="utf-8")
    (b / "notes.txt").write_text("från b", encoding="utf-8")
    out = tmp_path / "ut"
    subprocess.run(exe + [str(a), str(b), "-o", str(out)], capture_output=True)
    got = sorted(p.read_text(encoding="utf-8") for p in out.glob("notes*.md"))
    assert len(got) == 2 and "från a" in got[0] + got[1] and "från b" in got[0] + got[1]
    kb = a / "kb.md"
    kb.write_text("GAMMAL", encoding="utf-8")
    subprocess.run(exe + [str(a), "--combine", str(kb), "--force"], capture_output=True)
    assert "GAMMAL" not in kb.read_text(encoding="utf-8")


def test_repeated_edges_keep_first_occurrence():
    pages = {i: f"Patient: Tolvan 19121212-1212\nInnehåll {i}\nmer text\nSida {i}" for i in range(8)}
    out = pdf.strip_repeated_edges(pages)
    assert "Tolvan" in out[0] and "Tolvan" not in out[3]
    assert "Sida 0" in out[0] and "Sida 5" not in out[5]


# --------------------------------------------------------------------------- PDF
def test_digital_pdf_with_background_image_is_not_ocred():
    """WEG40120-mönstret: helsidesbild + SYNLIG text får aldrig OCR:as om."""
    doc = extract_file(FIX / "digital_bakgrund.pdf")
    assert {s["method"] for s in doc.sections} == {"layout"}
    assert "Besvärjelser kräver" in text(doc)
    assert len(doc.outline) == 3
    md = render_md(doc)
    assert "<!-- s. 2 -->" in md and "Kapitel 2 — s. 2" in md


@pytest.mark.parametrize("name", [
    "scan_ren.pdf", "scan_sned_ojamn.pdf", "scan_blek.pdf", "scan_skrivmaskin_90.pdf",
    "scan_upp_och_ner.pdf", "scan_fax_en.pdf", "foto_journal.jpg",
])
def test_scans_ocr_accuracy(fixtures, name):
    doc = extract_file(FIX / name)
    err = cer(text(doc), "\n".join(fixtures[name]))
    assert err < 0.03, f"{name}: CER {err:.1%}"
    assert "19121212-1212" in text(doc) or name == "scan_fax_en.pdf"


def test_corrupt_hidden_layer_is_reocred_and_blank_detected(fixtures):
    doc = extract_file(FIX / "scan_korrupt_lager_3s.pdf")
    methods = [s["method"] for s in doc.sections]
    assert methods == ["ocr", "blank", "ocr"]
    assert cer(text(doc), "\n".join(fixtures["scan_korrupt_lager_3s.pdf"])) < 0.03


def test_skew_estimate_on_uneven_background():
    """Grå rotationskant + gradient gav -6,5° i stället för -3,2° (2026-10-10)."""
    import io

    import pymupdf
    from PIL import Image

    from xt import preprocess as pp

    d = pymupdf.open(FIX / "scan_sned_ojamn.pdf")
    pix = d[0].get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
    g = pp.to_gray(Image.open(io.BytesIO(pix.tobytes("png"))))
    assert abs(pp.estimate_skew(g) - (-3.2)) < 0.4


def test_coverage_detects_dropped_lines():
    """Konfidens ser inte tappade rader – täckningen ska göra det."""
    import io

    import pymupdf
    from PIL import Image

    from xt import preprocess as pp

    d = pymupdf.open(FIX / "scan_ren.pdf")
    im = Image.open(io.BytesIO(d[0].get_pixmap(dpi=300, colorspace=pymupdf.csGRAY).tobytes("png")))
    res = ocr.ocr_robust(im, "swe+eng", "fast")
    assert res["coverage"] > 0.9
    norm = pp.normalize_illumination(pp.to_gray(im))
    blobs = pp.char_blobs(norm)
    ys = sorted({w["y"] for w in res["words"]})
    cut = ys[len(ys) // 2]
    half = [w for w in res["words"] if w["y"] < cut]
    assert pp.coverage(half, blobs) < 0.8


# --------------------------------------------------------------------------- Office m.m.
def test_docx():
    md = text(extract_file(FIX / "rapport.docx"))
    assert "# Rapport om knäskador" in md
    assert "## Åtgärder" in md
    assert "**fetstil**" in md and "*kursiv*" in md
    assert "[källan](https://example.se/knä)" in md
    assert re.search(r"^- Vila\n- Fysioterapi", md, re.M)
    assert re.search(r"^1\. Första steget\n2\. Andra steget", md, re.M)
    assert "a\\|b" in md
    assert "**normal**" not in md  # <w:b w:val="0"/> är INTE fetstil


def test_xlsx():
    md = text(extract_file(FIX / "budget.xlsx"))
    assert "## Blad: Kostnader" in md
    assert "| Hyra | 8450 | 2026-01-31 |" in md
    assert "612.5" in md
    assert "*(tomt blad)*" in md


def test_pptx():
    md = text(extract_file(FIX / "presentation.pptx"))
    assert "## Bild 1: Kvartalsöversikt" in md
    assert "- Intäkter ökade" in md
    assert "Talaranteckningar:** Säg något om Q3" in md
    assert "| Q3 | 1200 |" in md


def test_html():
    md = text(extract_file(FIX / "sida.html"))
    assert "# Rubrik ett" in md and "[länk](https://x.se)" in md
    assert "alert(" not in md and "Meny" not in md


def test_csv_cp1252():
    md = text(extract_file(FIX / "lista_cp1252.csv"))
    assert "| Åsa | Göteborg |" in md and "Malmö – syd" in md


def test_eml():
    doc = extract_file(FIX / "brev.eml")
    assert "Möte i morgon" in text(doc) and "kl. 10 på kontoret" in text(doc)


@pytest.mark.parametrize("name,needles", [
    ("rapport.doc", ["# Rapport om knäskador", "**fetstil**", "1. Första steget"]),
    ("rapport.odt", ["# Rapport om knäskador", "1. Första steget"]),
    ("budget.xls", ["| Hyra | 8450 | 2026-01-31 |"]),
    ("budget.ods", ["| Hyra | 8450 | 2026-01-31 |"]),
    ("presentation.ppt", ["Kvartalsöversikt", "Talaranteckningar:** Säg något om Q3"]),
])
def test_legacy_formats_via_libreoffice(name, needles):
    """Skapas av LibreOffice från de moderna fixturerna (se make_fixtures.build_legacy)."""
    if not (FIX / name).exists():
        pytest.skip(f"{name} saknas – kör make_fixtures.py med LibreOffice installerat")
    md = text(extract_file(FIX / name))
    for n in needles:
        assert n in md, f"{name}: saknar {n!r}"


# --------------------------------------------------------------------------- rendering
def test_txt_and_jsonl_render():
    doc = extract_file(FIX / "rapport.docx")
    t = render_txt(doc)
    assert "**" not in t and "Rapport om knäskador" in t
    recs = list(iter_jsonl(doc))
    assert recs and all("text" in r for r in recs)
    pdoc = extract_file(FIX / "digital_bakgrund.pdf")
    assert [r["page"] for r in iter_jsonl(pdoc)] == [1, 2, 3]


def test_cli_stdout_and_skip(tmp_path):
    exe = [sys.executable, str(SCRIPTS / "extract.py")]
    r = subprocess.run(exe + [str(FIX / "rapport.docx"), "--stdout"], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0 and "Rapport om knäskador" in r.stdout
    r1 = subprocess.run(exe + [str(FIX / "brev.eml"), "-o", str(tmp_path)], capture_output=True, text=True, encoding="utf-8")
    assert (tmp_path / "brev.md").exists() and r1.returncode == 0
    r2 = subprocess.run(exe + [str(FIX / "brev.eml"), "-o", str(tmp_path)], capture_output=True, text=True, encoding="utf-8")
    assert "[SKIP]" in r2.stderr


# --------------------------------------------------------------------------- structured
def test_structured_helpers():
    assert to_int("–2") == -2 and to_int("−3") == -3 and to_int("+4") == 4
    marker = re.compile(rf"^CREATURE ({SIGN}?\d+)$")
    lines = [(1, "ANIMATED BROOM"), (1, "CREATURE –1"), (1, "Perception +3"), (2, "AXIOMITE"), (2, "CREATURE 8"), (2, "text")]
    chunks = Book.chunk_by_pattern(lines, marker)
    assert [c["name"] for c in chunks] == ["ANIMATED BROOM", "AXIOMITE"]
    assert chunks[0]["match"].group(1) == "–1"
    body, cut = trim_chunk_body([(1, "CONCENTRATE"), (1, "Some prose here"), (1, "12"), (1, "after")])
    assert [l for _, l in body] == ["CONCENTRATE", "Some prose here"] and not cut
