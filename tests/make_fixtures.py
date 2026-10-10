"""Bygger testfiler med KÄNT facit i tests/fixtures/.

Inskannade "journaler" (fiktiva; testpersonnummer från Skatteverket) med typiska
skador: snedskanning, ojämn belysning, bleka kopior, fax, 90/180°-rotation,
korrupt dolt OCR-lager. Plus Office/HTML/e-post/CSV med kända detaljer.

    python tests/make_fixtures.py
"""
from __future__ import annotations

import io
import json
import random
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

HERE = Path(__file__).resolve().parent
FIX = HERE / "fixtures"
FIX.mkdir(exist_ok=True)
random.seed(7)
np.random.seed(7)

JOURNAL = [
    "Journalanteckning – Vårdcentralen Gustavsberg",
    "Patient: Tolvan Tolvansson, 19121212-1212",
    "Datum: 2026-03-14 kl. 10:32  Signerad av: Leg. läkare Åsa Öberg",
    "Kontaktorsak: Besvär från höger knä sedan cirka tre veckor.",
    "Anamnes: Patienten har tidigare varit väsentligen frisk. Inga kända",
    "allergier. Röker inte. Arbetar som snickare och går mycket i trappor.",
    "Smärtan förvärras vid knäböj och lindras av vila. Ingen svullnad.",
    "Status: Allmäntillstånd gott. Höger knä utan rodnad, ingen hydrops.",
    "Ömhet över mediala ledspringan. McMurray positivt medialt.",
    "Bedömning: Misstänkt menisklesion medialt höger knä.",
    "Åtgärd: Remiss till fysioterapeut. Ipren 400 mg vid behov, högst",
    "3 tabletter per dygn. Återbesök om 6 veckor eller tidigare vid behov.",
]
FAX_EN = [
    "REFERRAL LETTER - Department of Orthopaedics",
    "Patient: Test Testsson, born 1912-12-12",
    "Date: 14 March 2026. Reference no. 2026/0417-B",
    "We kindly ask you to assess this patient for a suspected medial",
    "meniscal tear of the right knee. Symptoms for three weeks, worse",
    "when squatting, no locking and no swelling. Please reply by fax.",
]


def font(name="arial.ttf", size=42):
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size)


def page(lines, fnt, size=(2480, 3508), margin=220, gap=1.75, bg=255, ink=0):
    img = Image.new("L", size, bg)
    d = ImageDraw.Draw(img)
    y = margin
    for i, ln in enumerate(lines):
        f = font(fnt.path.split("\\")[-1].split("/")[-1], int(fnt.size * 1.3)) if i == 0 else fnt
        d.text((margin, y), ln, font=f, fill=ink)
        y += int(f.size * gap) + (30 if i == 0 else 0)
    # en "stämpel" och en signaturkrumelur – vanligt brus i journaler
    d.ellipse((1700, 2700, 2150, 3000), outline=90, width=6)
    d.text((1760, 2810), "INSKANNAD", font=font("arialbd.ttf", 48), fill=90)
    pts = [(300 + k * 12, 2950 + int(25 * np.sin(k / 3))) for k in range(60)]
    d.line(pts, fill=40, width=4)
    return img


def uneven(img, strength=70):
    a = np.asarray(img).astype(np.float32)
    h, w = a.shape
    yy, xx = np.mgrid[0:h, 0:w]
    grad = strength * (xx / w) * 0.6 + strength * (yy / h) * 0.4
    return Image.fromarray(np.clip(a - grad, 0, 255).astype(np.uint8))


def noise(img, amount=18):
    a = np.asarray(img).astype(np.int16)
    a = a + np.random.normal(0, amount, a.shape).astype(np.int16)
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8))


def salt_pepper(img, p=0.004):
    a = np.asarray(img).copy()
    m = np.random.random(a.shape)
    a[m < p / 2] = 0
    a[m > 1 - p / 2] = 255
    return Image.fromarray(a)


def to_pdf_page(doc, img, dpi=300, hidden_text=None):
    import pymupdf

    w, h = img.size
    pw, ph = w * 72 / dpi, h * 72 / dpi
    pg = doc.new_page(width=pw, height=ph)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    pg.insert_image(pg.rect, stream=buf.getvalue())
    if hidden_text:  # osynligt (render mode 3) korrupt OCR-lager, som gamla skannrar gör
        y = 60
        for ln in hidden_text:
            pg.insert_text((50, y), ln, fontsize=9, render_mode=3)
            y += 14


def garble(s: str) -> str:
    tbl = str.maketrans({"t": "l", "e": "c", "a": "u", "i": "l", "o": "0", "s": "5", "r": "f", "n": "u", "ä": "a", "ö": "o"})
    return "".join(ch.translate(tbl) if random.random() < 0.45 else ch for ch in s).replace(" ", "" if random.random() < 0.3 else " ")


def build_scans():
    import pymupdf

    truth = {}
    f_arial, f_cour, f_times = font("arial.ttf"), font("cour.ttf", 44), font("times.ttf", 46)
    cases = {
        "ren": page(JOURNAL, f_arial),
        "sned_ojamn": noise(uneven(page(JOURNAL, f_times), 90).rotate(3.2, resample=Image.BICUBIC, fillcolor=235), 14),
        "blek": page(JOURNAL, f_arial, bg=232, ink=165).filter(ImageFilter.GaussianBlur(1.6)),
        "skrivmaskin_90": page(JOURNAL, f_cour).rotate(90, expand=True, fillcolor=255),
        "upp_och_ner": noise(page(JOURNAL, f_times).rotate(180), 10),
    }
    fax = page(FAX_EN, f_times)
    fax = fax.resize((fax.width // 3, fax.height // 3), Image.BILINEAR).point(lambda v: 0 if v < 150 else 255)
    fax = salt_pepper(fax.resize((fax.width * 3, fax.height * 3), Image.NEAREST).rotate(-1.4, fillcolor=255))
    cases["fax_en"] = fax

    for name, img in cases.items():
        doc = pymupdf.open()
        to_pdf_page(doc, img)
        doc.save(FIX / f"scan_{name}.pdf", garbage=3, deflate=True)
        truth[f"scan_{name}.pdf"] = FAX_EN if name == "fax_en" else JOURNAL

    # flersidig: korrupt dolt OCR-lager + en tom sida + en ren sida
    doc = pymupdf.open()
    to_pdf_page(doc, noise(uneven(page(JOURNAL, f_arial), 50), 10), hidden_text=[garble(l) for l in JOURNAL] * 4)
    to_pdf_page(doc, Image.new("L", (2480, 3508), 250))
    to_pdf_page(doc, page(JOURNAL[:6], f_cour))
    doc.save(FIX / "scan_korrupt_lager_3s.pdf", garbage=3, deflate=True)
    truth["scan_korrupt_lager_3s.pdf"] = JOURNAL + JOURNAL[:6]

    # mobilfoto av papper (JPEG, grått, snett, lite oskärpa)
    img = noise(uneven(page(JOURNAL, f_arial), 60).rotate(-2.1, resample=Image.BICUBIC, fillcolor=200), 12)
    img = img.filter(ImageFilter.GaussianBlur(0.8)).resize((1240, 1754))
    img.convert("RGB").save(FIX / "foto_journal.jpg", quality=72)
    truth["foto_journal.jpg"] = JOURNAL
    (FIX / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=1), encoding="utf-8")


def build_digital():
    """Digital PDF med riktigt textlager + bakgrundsbild över hela sidan
    (fallet som INTE får OCR:as om – WEG40120-mönstret)."""
    import pymupdf

    doc = pymupdf.open()
    for n in range(1, 4):
        pg = doc.new_page()
        bg = Image.new("RGB", (600, 800), (238, 228, 205))
        buf = io.BytesIO()
        bg.save(buf, "PNG")
        pg.insert_image(pg.rect, stream=buf.getvalue())
        pg.insert_text((72, 60), "Mystik och magi – kapitel " + str(n), fontsize=20)
        body = ("Besvärjelser kräver att magikern har tillräcklig kännedom om sin aspekt. "
                "Den som misslyckas med slaget förlorar ändå sin koncentration och måste vila. ") * 6
        pg.insert_textbox(pymupdf.Rect(72, 100, 520, 700), body, fontsize=11)
        pg.insert_text((290, 800), str(n), fontsize=9)
    doc.set_toc([[1, "Kapitel 1", 1], [1, "Kapitel 2", 2], [1, "Kapitel 3", 3]])
    doc.save(FIX / "digital_bakgrund.pdf", garbage=3, deflate=True)


def build_office():
    import docx
    import openpyxl
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.util import Inches

    d = docx.Document()
    d.add_heading("Rapport om knäskador", 1)
    p = d.add_paragraph("Detta är ")
    p.add_run("fetstil").bold = True
    p.add_run(" och ")
    p.add_run("kursiv").italic = True
    p.add_run(". Se ")
    # hyperlänk
    rid = d.part.relate_to("https://example.se/knä", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "källan"
    r.append(t)
    hl.append(r)
    p._p.append(hl)
    d.add_heading("Åtgärder", 2)
    d.add_paragraph("Vila", style="List Bullet")
    d.add_paragraph("Fysioterapi", style="List Bullet")
    d.add_paragraph("Första steget", style="List Number")
    d.add_paragraph("Andra steget", style="List Number")
    tb = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["Datum", "Värde", "Not"], ["2026-03-14", "42", "a|b"], ["2026-04-01", "17", "ök"]]):
        for j, v in enumerate(row):
            tb.cell(i, j).text = v
    np_ = d.add_paragraph("Ej fet: ")
    rr = np_.add_run("normal")
    b = OxmlElement("w:b")
    b.set(qn("w:val"), "0")
    rr._r.get_or_add_rPr().append(b)
    d.save(FIX / "rapport.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kostnader"
    ws.append(["Post", "Belopp", "Datum"])
    import datetime as dt
    ws.append(["Hyra", 8450, dt.datetime(2026, 1, 31)])
    ws.append(["El", 612.5, dt.datetime(2026, 2, 28)])
    ws.append(["Summa", "=SUM(B2:B3)", None])
    wb.create_sheet("Tomt")
    wb.save(FIX / "budget.xlsx")

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Kvartalsöversikt"
    s.placeholders[1].text = "Intäkter ökade\nKostnader minskade"
    s.notes_slide.notes_text_frame.text = "Säg något om Q3 här."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Tabell"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text, tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Q", "Kr", "Q3", "1200"
    prs.save(FIX / "presentation.pptx")

    (FIX / "sida.html").write_text(
        "<html><head><title>Testsida</title><script>alert(1)</script></head><body><nav>Meny</nav><main>"
        "<h1>Rubrik ett</h1><p>Text med <a href='https://x.se'>länk</a> och <b>fetstil</b>.</p>"
        "<ul><li>Ett</li><li>Två</li></ul><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "</main></body></html>", encoding="utf-8")
    (FIX / "lista_cp1252.csv").write_bytes("namn;stad\nÅsa;Göteborg\nÖrjan;Malmö – syd\n".encode("cp1252"))
    (FIX / "brev.eml").write_bytes(
        "From: Åsa <asa@example.se>\nTo: test@example.se\nSubject: =?utf-8?q?M=C3=B6te_i_morgon?=\n"
        "MIME-Version: 1.0\nContent-Type: text/plain; charset=utf-8\nContent-Transfer-Encoding: 8bit\n\n"
        "Hej! Vi ses kl. 10 på kontoret.\n".encode("utf-8"))


def _art(size=(900, 650)):
    """En "illustration": konturer, skraffering och en tät figur – som en tecknad bild."""
    img = Image.new("L", size, 255)
    d = ImageDraw.Draw(img)
    w, h = size
    d.rectangle((4, 4, w - 5, h - 5), outline=0, width=6)
    for k in range(0, w, 14):
        d.line((k, h * 0.65, k + 60, h - 10), fill=40, width=3)
    d.ellipse((w * 0.3, h * 0.12, w * 0.7, h * 0.6), outline=0, width=8, fill=170)
    d.polygon([(w * 0.1, h * 0.9), (w * 0.25, h * 0.3), (w * 0.4, h * 0.9)], outline=0, fill=90)
    for k in range(40):
        d.line((w * 0.55 + k * 6, h * 0.65, w * 0.6 + k * 6, h * 0.95), fill=0, width=2)
    return img


def build_figures():
    import pymupdf

    # 1) Digital PDF: vektorkarta med bildtext + korsreferenser + tryckta sidnummer
    doc = pymupdf.open()
    for n in range(1, 6):
        pg = doc.new_page()
        pg.insert_text((72, 60), f"Kapitel {n}", fontsize=18)
        body = ("Stationen ligger i ett asteroidbälte. Se Figur 1 för däcksplanen och läs mer på sidan 4. " * 3
                if n == 2 else "Vanlig brödtext om stationen och dess besättning. " * 12)
        pg.insert_textbox(pymupdf.Rect(72, 80, 520, 200), body, fontsize=10)
        if n == 1:
            for k in range(40):
                pg.draw_line((100 + k * 8, 230), (100 + k * 5, 470), color=(0, 0, 0.5), width=1.5)
            pg.draw_circle((300, 350), 90, color=(0.7, 0, 0), fill=(1, 0.9, 0.8))
            pg.draw_rect((95, 225, 500, 480), color=(0, 0, 0), width=2)
            pg.insert_text((150, 498), "Figur 1: Däcksplan för stationen", fontsize=9)
        # samma logotyp på varje sida -> ska filtreras bort som dekor
        logo = Image.new("RGB", (120, 120), (20, 60, 160))
        ImageDraw.Draw(logo).ellipse((10, 10, 110, 110), fill=(240, 200, 0))
        buf = io.BytesIO()
        logo.save(buf, "PNG")
        pg.insert_image(pymupdf.Rect(480, 700, 560, 780), stream=buf.getvalue())
        pg.insert_text((290, 815), str(n + 2), fontsize=9)  # tryckt sidnr = PDF-sida + 2
    doc.save(FIX / "figurer_digital.pdf", garbage=3, deflate=True)

    # 2) Inskannad sida: text, illustration, bildtext under bilden
    page = Image.new("L", (2480, 3508), 255)
    d = ImageDraw.Draw(page)
    f = font("times.ttf", 46)
    y = 220
    for ln in ["The smuggler's ship was a battered freighter.", "It had seen better days, but it was fast",
               "enough to outrun most patrol craft in the sector."] * 2:
        d.text((220, y), ln, font=f, fill=0)
        y += 80
    page.paste(_art((1700, 1200)), (380, 760))
    d.text((380, 1990), "Figure 3: The freighter at Mos Eisley.", font=font("timesi.ttf", 38), fill=0)
    y = 2150
    for ln in ["Below the picture the text continues with more", "details about the ship and its owner.",
               "The captain never trusted anyone on board."] * 2:
        d.text((220, y), ln, font=f, fill=0)
        y += 80
    doc = pymupdf.open()
    to_pdf_page(doc, noise(page, 6))
    doc.save(FIX / "figurer_scan.pdf", garbage=3, deflate=True)

    # 3) Word med inbäddad bild + bildtext (formatmall Caption)
    import docx
    from docx.shared import Inches

    art = _art((600, 420)).convert("RGB")
    ap = FIX / "_art.png"
    art.save(ap)
    dd = docx.Document()
    dd.add_paragraph("Text före bilden.")
    dd.add_picture(str(ap), width=Inches(3))
    dd.add_paragraph("Figur 2: Skiss av stationen", style="Caption")
    dd.add_paragraph("Text efter bilden, se Figur 2.")
    dd.save(FIX / "rapport_bild.docx")

    # 4) PowerPoint med bild + alt-text
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Bildbild"
    pic = s.shapes.add_picture(str(ap), Inches(1), Inches(2), Inches(4))
    pic._element.nvPicPr.cNvPr.set("descr", "Skiss av stationen")
    prs.save(FIX / "presentation_bild.pptx")
    ap.unlink()


def build_legacy():
    """Äldre/öppna format via LibreOffice (hoppas över om det saknas)."""
    import subprocess
    import tempfile

    so = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
    if not so.exists():
        print("  (LibreOffice saknas – hoppar över .doc/.odt/.xls/.ods/.ppt)")
        return
    prof = Path(tempfile.mkdtemp(prefix="xt_fix_")).as_uri()
    for src, fmt in [("rapport.docx", "doc"), ("rapport.docx", "odt"), ("budget.xlsx", "xls"),
                     ("budget.xlsx", "ods"), ("presentation.pptx", "ppt")]:
        subprocess.run([str(so), f"-env:UserInstallation={prof}", "--headless", "--convert-to", fmt,
                        "--outdir", str(FIX), str(FIX / src)], capture_output=True, timeout=240)


if __name__ == "__main__":
    build_scans()
    build_digital()
    build_office()
    build_figures()
    build_legacy()
    print("Fixtures i", FIX)
    for f in sorted(FIX.iterdir()):
        print(f"  {f.name:32} {f.stat().st_size:>9,}")
