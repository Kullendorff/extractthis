"""Office-format -> markdown: DOCX, XLSX/XLS, PPTX + äldre/öppna format via
konvertering (LibreOffice headless, Word/PowerPoint-COM som reserv)."""
from __future__ import annotations

import datetime as dt
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from .common import SOFFICE_EXE, log, md_table

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_CAP = "\x00CAP\x00"  # intern markör för bildtextstycken (Caption/Beskrivning)
_IMG_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
           "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
           "v": "urn:schemas-microsoft-com:vml"}


# ---------------------------------------------------------------------------
# Kryptering
# ---------------------------------------------------------------------------

def open_maybe_encrypted(path: Path, password: str | None):
    """Returnerar sökväg eller BytesIO (dekrypterad). Kastar PermissionError om
    filen är krypterad och lösenord saknas/är fel."""
    try:
        import msoffcrypto
    except ImportError:
        return path
    try:
        with open(path, "rb") as f:
            of = msoffcrypto.OfficeFile(f)
            if not of.is_encrypted():
                return path
            if not password:
                raise PermissionError("Filen är lösenordsskyddad (ange --password)")
            of.load_key(password=password)
            buf = io.BytesIO()
            of.decrypt(buf)
            buf.seek(0)
            return buf
    except PermissionError:
        raise
    except Exception:
        return path  # inte ett OLE/OOXML-format msoffcrypto känner igen


# ---------------------------------------------------------------------------
# Konvertering av äldre format
# ---------------------------------------------------------------------------

def _kill_tree(proc: subprocess.Popen) -> None:
    """Vid timeout räcker inte proc.kill(): soffice.exe är bara en startare och
    soffice.bin lever kvar. Döda hela processträdet."""
    import os
    import signal

    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(proc.pid)], capture_output=True, timeout=30)
        else:
            os.killpg(proc.pid, signal.SIGKILL)
    except Exception:
        proc.kill()


def _convert_with_soffice(path: Path, target: str, tmp: Path, timeout: int = 240) -> Path | None:
    """LibreOffice headless med EGEN profil – annars hänger den om Johan har
    LibreOffice öppet, eller om profilen är låst."""
    import os

    if not SOFFICE_EXE:
        return None
    profile = (tmp / "profile").as_uri()
    # absolut sökväg: en fil som heter "-x.doc" får inte tolkas som en flagga
    cmd = [SOFFICE_EXE, f"-env:UserInstallation={profile}", "--headless", "--norestore",
           "--convert-to", target, "--outdir", str(tmp), str(Path(path).resolve())]
    kw = {"start_new_session": True} if os.name != "nt" else {}
    proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **kw)
    try:
        proc.wait(timeout=timeout)
    except subprocess.TimeoutExpired:
        log(f"    [!] LibreOffice svarade inte på {timeout}s – avbryter")
        _kill_tree(proc)
        return None
    out = tmp / (Path(path).stem + "." + target.split(":")[0])
    return out if out.exists() else None


def _convert_with_office_com(path: Path, target: str, tmp: Path) -> Path | None:
    """Word/PowerPoint/Excel via COM (bäst trohet för .doc/.ppt, men kan hänga på
    dialogrutor – därför bara reserv). Dummylösenord gör att Open KASTAR för
    lösenordsskyddade filer i stället för att visa en modal dialog som hänger."""
    try:
        import win32com.client
    except ImportError:
        return None
    out = tmp / (Path(path).stem + "." + target)
    src = str(Path(path).resolve())
    app = None
    try:
        if target == "docx":
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            d = app.Documents.Open(src, ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False,
                                   PasswordDocument="xt-ingen", WritePasswordDocument="xt-ingen")
            d.SaveAs2(str(out), FileFormat=16)
            d.Close(False)
        elif target == "pptx":
            app = win32com.client.DispatchEx("PowerPoint.Application")
            p = app.Presentations.Open(src + "::xt-ingen::", ReadOnly=True, WithWindow=False)
            p.SaveAs(str(out), 24)
            p.Close()
        elif target == "xlsx":
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            wb = app.Workbooks.Open(src, ReadOnly=True, Password="xt-ingen", WriteResPassword="xt-ingen")
            wb.SaveAs(str(out), 51)
            wb.Close(False)
    except Exception as e:
        log(f"    [!] Office-COM: {e}")
        return None
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
    return out if out.exists() else None


def convert_legacy(path: Path, target: str) -> io.BytesIO:
    """Konvertera till target (docx/pptx/xlsx) och returnera resultatet i minnet.
    Temp-mappen (inkl. en 10–30 MB LibreOffice-profil) raderas alltid – annars
    fyller en rekursiv körning över hundratals .doc-filer %TEMP%."""
    tmp = Path(tempfile.mkdtemp(prefix="xt_conv_"))
    try:
        out = _convert_with_soffice(path, target, tmp) or _convert_with_office_com(path, target, tmp)
        if not out:
            raise RuntimeError(f"Kunde inte konvertera {Path(path).suffix} -> {target} "
                               "(varken LibreOffice eller Office fungerade)")
        buf = io.BytesIO(out.read_bytes())
        buf.name = out.name  # så att filändelsen kan läsas av mottagaren
        return buf
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _suffix(src) -> str:
    return Path(getattr(src, "name", None) or str(src)).suffix.lower()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _on(el) -> bool:
    """<w:b/> är på, <w:b w:val="0"/> eller "false" är AV."""
    if el is None:
        return False
    v = el.get(W + "val")
    return v not in ("0", "false", "off")


class _Docx:
    def __init__(self, src, images_dir: Path | None, stem: str):
        import docx

        self.doc = docx.Document(src)
        self.rels = {rid: r for rid, r in self.doc.part.rels.items()}
        self.style_names = {s.style_id: (s.name or "") for s in self.doc.styles}
        # Numrering kan ärvas från stilen ("List Number") i stället för att stå på stycket
        self.style_num = {}
        for s in self.doc.styles:
            ppr = s.element.find(W + "pPr")
            np_ = ppr.find(W + "numPr") if ppr is not None else None
            if np_ is not None:
                self.style_num[s.style_id] = np_
        self.numfmt = self._load_numbering()
        self.counters: dict = {}
        self.images_dir, self.stem = images_dir, stem
        self.image_files: list[str] = []
        self.fig_dir: Path | None = None
        self.figs: list[dict] = []
        self._fig_by_rid: dict[str, str] = {}

    def _figure(self, rid: str) -> str | None:
        """Spara en inbäddad bild (en gång per relation) och returnera platshållaren."""
        from .figures import fig_id, placeholder

        if rid in self._fig_by_rid:
            return placeholder(self._fig_by_rid[rid])
        rel = self.rels.get(rid)
        if rel is None or rel.is_external or "image" not in rel.reltype:
            return None
        part = rel.target_part
        ext = Path(str(part.partname)).suffix.lower() or ".png"
        if ext in (".emf", ".wmf"):  # vektor-metafiler kan inte visas i md/Obsidian
            return None
        k = len(self.figs) + 1
        fid = fig_id(None, k)
        self.fig_dir.mkdir(parents=True, exist_ok=True)
        fn = self.fig_dir / f"{self.stem}_{k:03d}{ext}"
        fn.write_bytes(part.blob)
        w = h = None
        try:
            from PIL import Image

            with Image.open(fn) as im:
                w, h = im.size
        except Exception:
            pass
        self.figs.append({"id": fid, "page": None, "file": str(fn), "caption": "", "source": "docx", "w": w, "h": h})
        self._fig_by_rid[rid] = fid
        return placeholder(fid)

    def _load_numbering(self) -> dict:
        """(numId, ilvl) -> numFmt ('decimal', 'bullet', ...)."""
        out = {}
        try:
            numbering = self.doc.part.numbering_part.element
        except Exception:
            return out
        abstract = {}
        for an in numbering.findall(W + "abstractNum"):
            aid = an.get(W + "abstractNumId")
            for lvl in an.findall(W + "lvl"):
                fmt = lvl.find(W + "numFmt")
                abstract[(aid, lvl.get(W + "ilvl"))] = fmt.get(W + "val") if fmt is not None else "bullet"
        for num in numbering.findall(W + "num"):
            nid = num.get(W + "numId")
            a = num.find(W + "abstractNumId")
            if a is None:
                continue
            aid = a.get(W + "val")
            for (aa, ilvl), fmt in abstract.items():
                if aa == aid:
                    out[(nid, ilvl)] = fmt
        return out

    # -- runs -----------------------------------------------------------
    def _runs(self, el, link=None):
        for ch in el:
            tag = ch.tag
            if tag == W + "r":
                rpr = ch.find(W + "rPr")
                b = i = code = False
                if rpr is not None:
                    b = _on(rpr.find(W + "b"))
                    i = _on(rpr.find(W + "i"))
                    f = rpr.find(W + "rFonts")
                    if f is not None and re.search(r"mono|courier|consolas", f.get(W + "ascii", "") or "", re.I):
                        code = True
                txt = []
                for t in ch:
                    if t.tag == W + "t":
                        txt.append(t.text or "")
                    elif t.tag == W + "tab":
                        txt.append("\t")
                    elif t.tag in (W + "br", W + "cr"):
                        txt.append("\n")
                    elif t.tag == W + "noBreakHyphen":
                        txt.append("-")
                if txt:
                    yield "".join(txt), b, i, code, link
                # Inbäddad bild i löptexten (DrawingML a:blip eller gammal VML v:imagedata)
                if self.fig_dir is not None:
                    # lxml direkt – python-docx:s egen xpath() tar inte namespaces=
                    from lxml import etree

                    for rid in etree._Element.xpath(ch, ".//a:blip/@r:embed | .//v:imagedata/@r:id",
                                                    namespaces=_IMG_NS):
                        tok = self._figure(rid)
                        if tok:
                            yield tok, False, False, False, None
            elif tag == W + "hyperlink":
                rid = ch.get(R_ID)
                url = None
                if rid and rid in self.rels and self.rels[rid].is_external:
                    url = self.rels[rid].target_ref
                yield from self._runs(ch, url)
            elif tag in (W + "ins", W + "smartTag", W + "fldSimple", W + "customXml",
                         W + "moveTo", W + "dir", W + "bdo"):  # moveFrom = gammal plats, hoppas över
                yield from self._runs(ch, link)
            elif tag == W + "sdt":
                c = ch.find(W + "sdtContent")
                if c is not None:
                    yield from self._runs(c, link)
            # w:del (borttagen text vid spårade ändringar) hoppas medvetet över

    def _inline(self, p_el) -> str:
        merged: list[list] = []
        for txt, b, i, code, link in self._runs(p_el):
            if merged and merged[-1][1:] == [b, i, code, link]:
                merged[-1][0] += txt
            else:
                merged.append([txt, b, i, code, link])
        out = []
        for txt, b, i, code, link in merged:
            core = txt.strip()
            if not core:
                out.append(txt)
                continue
            lead, trail = txt[: len(txt) - len(txt.lstrip())], txt[len(txt.rstrip()):]
            if code:
                core = f"`{core}`"
            elif b and i:
                core = f"***{core}***"
            elif b:
                core = f"**{core}**"
            elif i:
                core = f"*{core}*"
            if link:
                core = f"[{core}]({link})"
            out.append(lead + core + trail)
        return "".join(out)

    # -- block ----------------------------------------------------------
    def _heading_level(self, p_el, style: str) -> int:
        m = re.search(r"(?:heading|rubrik)\s*(\d)", style, re.I)
        if m:
            return min(int(m.group(1)), 6)
        if style.lower() in ("title", "rubrik", "titel"):
            return 1
        if style.lower() in ("subtitle", "underrubrik", "undertitel"):
            return 2
        ppr = p_el.find(W + "pPr")
        if ppr is not None:
            ol = ppr.find(W + "outlineLvl")
            if ol is not None and ol.get(W + "val", "9").isdigit() and int(ol.get(W + "val")) < 6:
                return int(ol.get(W + "val")) + 1
        return 0

    def paragraph(self, p_el) -> str:
        text = self._inline(p_el).strip()
        if not text:
            return ""
        ppr = p_el.find(W + "pPr")
        style = ""
        num = None
        if ppr is not None:
            ps = ppr.find(W + "pStyle")
            if ps is not None:
                sid = ps.get(W + "val", "")
                style = self.style_names.get(sid, sid)
                num = self.style_num.get(sid)
            num = ppr.find(W + "numPr") if ppr.find(W + "numPr") is not None else num
        lvl = self._heading_level(p_el, style)
        if lvl:
            return "#" * lvl + " " + re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        if num is not None:
            il = num.find(W + "ilvl")
            ni = num.find(W + "numId")
            ilvl = il.get(W + "val", "0") if il is not None else "0"
            nid = ni.get(W + "val", "") if ni is not None else ""
            fmt = self.numfmt.get((nid, ilvl), "bullet")
            indent = "  " * int(ilvl or 0)
            if fmt in ("decimal", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman"):
                key = (nid, ilvl)
                self.counters[key] = self.counters.get(key, 0) + 1
                return f"{indent}{self.counters[key]}. {text}"
            return f"{indent}- {text}"
        if re.search(r"list|punkt", style, re.I):
            return "- " + text
        if re.search(r"quote|citat", style, re.I):
            return "> " + text
        if re.search(r"caption|beskrivning|bildtext", style, re.I):
            return _CAP + text  # kopplas till närmast föregående bild i to_markdown
        return text

    def table(self, tbl_el) -> str:
        rows = []
        for tr in tbl_el.findall(W + "tr"):
            row = []
            for tc in tr.findall(W + "tc"):
                tcpr = tc.find(W + "tcPr")
                span = 1
                cont = False
                if tcpr is not None:
                    gs = tcpr.find(W + "gridSpan")
                    if gs is not None:
                        span = int(gs.get(W + "val", "1"))
                    vm = tcpr.find(W + "vMerge")
                    cont = vm is not None and vm.get(W + "val") in (None, "continue")
                paras = [self._inline(p).strip() for p in tc.iter(W + "p")]
                cell = "<br>".join(p for p in paras if p)
                row.append("" if cont else cell)
                row.extend([""] * (span - 1))
            rows.append(row)
        return md_table(rows)

    @staticmethod
    def _in_fallback(el) -> bool:
        """Textrutor finns ofta dubbelt: mc:Choice (DrawingML) + mc:Fallback (VML)."""
        p = el.getparent()
        while p is not None:
            if p.tag.endswith("}Fallback"):
                return True
            p = p.getparent()
        return False

    def blocks(self, parent):
        for el in parent:
            if el.tag == W + "p":
                yield self.paragraph(el)
                # Textrutor (flygblad, formulär) ligger inbäddade i stycket och
                # syns inte för _runs – läs dem som egna block efter stycket.
                for tb in el.iter(W + "txbxContent"):
                    if not self._in_fallback(tb):
                        yield from self.blocks(tb)
            elif el.tag == W + "tbl":
                yield self.table(el)
            elif el.tag == W + "sdt":
                c = el.find(W + "sdtContent")
                if c is not None:
                    yield from self.blocks(c)

    def images(self) -> list[str]:
        if self.images_dir is None:
            return []
        files = []
        for rel in self.rels.values():
            if "image" in rel.reltype and not rel.is_external:
                part = rel.target_part
                name = f"{self.stem}_{Path(part.partname).name}"
                self.images_dir.mkdir(parents=True, exist_ok=True)
                (self.images_dir / name).write_bytes(part.blob)
                files.append(name)
        return files

    def notes(self) -> str:
        """Fotnoter och slutnoter (separatorposter med id -1/0 hoppas över)."""
        from docx.oxml import parse_xml

        out = []
        for part in self.doc.part.package.parts:
            name = str(part.partname)
            if name not in ("/word/footnotes.xml", "/word/endnotes.xml"):
                continue
            root = parse_xml(part.blob)
            tag = "footnote" if "foot" in name else "endnote"
            items = []
            for fn in root.findall(W + tag):
                fid = fn.get(W + "id", "")
                if fn.get(W + "type") in ("separator", "continuationSeparator", "continuationNotice"):
                    continue
                txt = " ".join(self._inline(p).strip() for p in fn.iter(W + "p")).strip()
                if txt:
                    items.append(f"[^{fid}]: {txt}")
            if items:
                out.append(("## Fotnoter" if tag == "footnote" else "## Slutnoter") + "\n\n" + "\n".join(items))
        return "\n\n".join(out)

    def to_markdown(self) -> str:
        from .figures import FIG_RE

        raw = [b for b in self.blocks(self.doc.element.body) if b and b.strip()]
        parts: list[str] = []
        for b in raw:
            if b.startswith(_CAP):
                cap = b[len(_CAP):].strip()
                prev = FIG_RE.findall(parts[-1]) if parts else []
                fig = next((f for f in self.figs if prev and f["id"] == prev[-1]), None)
                if fig is not None and not fig["caption"]:
                    fig["caption"] = cap  # bildtexten följer med bilden i stället för som löptext
                    continue
                b = cap
            parts.append(b)
        notes = self.notes()
        if notes:
            parts.append(notes)
        # slå ihop följande listrader till en lista (enkelradbrytning)
        out: list[str] = []
        for b in parts:
            if out and re.match(r"^\s*(-|\d+\.) ", b) and re.match(r"^\s*(-|\d+\.) ", out[-1].splitlines()[-1]):
                out[-1] += "\n" + b
            else:
                out.append(b)
        return "\n\n".join(out)


def extract_docx(path, images_dir: Path | None = None, password: str | None = None,
                 stem: str | None = None, figures_dir: Path | None = None) -> tuple[str, dict]:
    src = open_maybe_encrypted(path, password) if isinstance(path, (str, Path)) else path
    d = _Docx(src, images_dir, stem or Path(getattr(path, "name", None) or str(path)).stem)
    d.fig_dir = figures_dir
    md = d.to_markdown()
    imgs = d.images()
    cp = d.doc.core_properties
    meta = {"title": cp.title or None, "author": cp.author or None, "images": imgs, "figures": d.figs}
    return md, meta


# ---------------------------------------------------------------------------
# Kalkylblad
# ---------------------------------------------------------------------------

def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _trim(rows: list[list]) -> list[list]:
    rows = [r for r in rows]
    while rows and not any(c != "" for c in rows[-1]):
        rows.pop()
    if not rows:
        return rows
    width = max((max((j + 1 for j, c in enumerate(r) if c != ""), default=0) for r in rows), default=0)
    rows = [r[:width] for r in rows]
    while rows and not any(c != "" for c in rows[0]):
        rows.pop(0)
    return rows


def _sheets_to_md(sheets, max_rows: int) -> tuple[str, dict]:
    """sheets: [(namn, rader (max max_rows+1 st), antal_rader_totalt)]."""
    parts, info = [], {}
    for name, rows, total in sheets:
        rows = _trim(rows)
        info[name] = total
        parts.append(f"## Blad: {name}")
        if not rows:
            parts.append("*(tomt blad)*")
            continue
        parts.append(md_table(rows[: max_rows + 1]))
        rest = total - min(total, len(rows))
        if rest > 0:
            parts.append(f"*… {rest:,} rader till utelämnade (öka --max-rows)*")
    return "\n\n".join(parts), {"sheets": info}


def _capped(rows_iter, max_rows: int):
    """Behåll bara de första max_rows+1 icke-tomma raderna, räkna resten –
    ett blad med en miljon rader ska inte läsas in i minnet."""
    keep, total = [], 0
    for row in rows_iter:
        if not any(c != "" for c in row):
            continue  # tomrader räknas inte (formaterade tomrader gav falskt "990 rader utelämnade")
        total += 1
        if len(keep) <= max_rows:
            keep.append(row)
    return keep, total


def read_sheets(path, password: str | None = None, max_rows: int = 2000):
    ext = _suffix(path)
    if ext in (".xlsx", ".xlsm", ".xltx"):
        import openpyxl

        src = open_maybe_encrypted(path, password) if isinstance(path, (str, Path)) else path
        # data_only=True: cachade formelvärden, inte "=SUMMA(A1:A9)"
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        try:
            out = []
            for ws in wb.worksheets:
                rows, total = _capped(([_fmt(c) for c in row] for row in ws.iter_rows(values_only=True)), max_rows)
                out.append((ws.title, rows, total))
            return out
        finally:
            wb.close()
    if ext == ".xls":
        import xlrd

        wb = xlrd.open_workbook(path, on_demand=True)
        try:
            return _xls_sheets(wb, xlrd, max_rows)
        finally:
            wb.release_resources()
    # .ods m.fl.
    return read_sheets(convert_legacy(path, "xlsx"), max_rows=max_rows)


def _xls_rows(sh, wb, xlrd):
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            v = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                except Exception:
                    pass
            row.append(_fmt(v))
        yield row


def _xls_sheets(wb, xlrd, max_rows):
    out = []
    for sh in wb.sheets():
        rows, total = _capped(_xls_rows(sh, wb, xlrd), max_rows)
        out.append((sh.name, rows, total))
    return out


def extract_sheet(path: Path, max_rows: int = 2000, password: str | None = None) -> tuple[str, dict]:
    return _sheets_to_md(read_sheets(path, password, max_rows), max_rows)


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _pptx_picture(shape, ctx: dict) -> str | None:
    """Spara bilden ur en bildform och returnera platshållaren (bildtext = alt-text)."""
    from .figures import fig_id, placeholder

    try:
        img = shape.image
    except Exception:
        return None
    if img.ext in ("emf", "wmf"):
        return None
    ctx["k"] = ctx.get("k", 0) + 1
    fid = fig_id(ctx["slide"], ctx["k"])
    ctx["fig_dir"].mkdir(parents=True, exist_ok=True)
    fn = ctx["fig_dir"] / f"{ctx['stem']}_bild{ctx['slide']:02d}_{ctx['k']:02d}.{img.ext}"
    fn.write_bytes(img.blob)
    descr = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
    w, h = (img.size if hasattr(img, "size") else (None, None))
    ctx["figs"].append({"id": fid, "page": ctx["slide"], "file": str(fn), "source": "pptx", "w": w, "h": h,
                        "caption": descr[0].strip() if descr and descr[0].strip() else ""})
    return placeholder(fid)


def _shape_md(shape, out: list[str], ctx: dict | None = None):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in sorted(shape.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            _shape_md(s, out, ctx)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        out.append(md_table([[c.text.strip() for c in r.cells] for r in shape.table.rows]))
        return
    if getattr(shape, "has_chart", False) and shape.has_chart:
        ch = shape.chart
        title = ch.chart_title.text_frame.text if ch.has_title else "diagram"
        try:
            cats = list(ch.plots[0].categories)
            rows = [[s.name] + [_fmt(v) for v in s.values] for s in ch.plots[0].series]
            out.append(f"*Diagram: {title}*\n\n" + md_table(rows, [""] + [str(c) for c in cats]))
        except Exception:
            out.append(f"*Diagram: {title}*")
        return
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        tok = _pptx_picture(shape, ctx) if ctx and ctx.get("fig_dir") is not None else None
        if tok:
            out.append(tok)
            return
        descr = shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
        if descr and descr[0].strip():
            out.append(f"*[Bild: {descr[0].strip()}]*")
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines = []
        for para in shape.text_frame.paragraphs:
            t = "".join(r.text for r in para.runs).strip() or para.text.strip()
            if t:
                lines.append(("  " * para.level + "- " + t) if (para.level or len(shape.text_frame.paragraphs) > 1) else t)
        if lines:
            out.append("\n".join(lines))


def extract_pptx(path: Path, password: str | None = None, figures_dir: Path | None = None,
                 stem: str | None = None) -> tuple[str, dict]:
    from pptx import Presentation

    src = open_maybe_encrypted(path, password) if isinstance(path, (str, Path)) else path
    prs = Presentation(src)
    stem = stem or Path(getattr(path, "name", None) or str(path)).stem
    figs: list[dict] = []
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        ctx = {"fig_dir": figures_dir, "stem": stem, "slide": n, "figs": figs}
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        parts.append(f"## Bild {n}" + (f": {title}" if title else ""))
        body: list[str] = []
        for s in sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            if title_shape is not None and s.shape_id == title_shape.shape_id:
                continue
            _shape_md(s, body, ctx)
        parts.extend(body)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("> **Talaranteckningar:** " + notes.replace("\n", "\n> "))
    return "\n\n".join(parts), {"slides": len(prs.slides), "figures": figs}


# ---------------------------------------------------------------------------
# Router för Office-familjen
# ---------------------------------------------------------------------------

_MAIN_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
}
_VARIANT_RE = {
    "docx": re.compile(r"application/vnd\.(?:ms-word\.(?:document|template)\.macroEnabledTemplate|"
                       r"ms-word\.document\.macroEnabled|openxmlformats-officedocument\.wordprocessingml\.template)"
                       r"\.main\+xml"),
    "pptx": re.compile(r"application/vnd\.(?:ms-powerpoint\.(?:presentation|slideshow|template)\.macroEnabled|"
                       r"openxmlformats-officedocument\.presentationml\.(?:slideshow|template))\.main\+xml"),
}


def normalize_ooxml(path: Path, kind: str) -> io.BytesIO:
    """.docm/.dotx/.pptm/.ppsx/.potx skiljer sig från .docx/.pptx bara i huvud-
    delens innehållstyp, men python-docx/pptx vägrar öppna dem. Skriv om typen i
    en kopia i minnet (makron körs aldrig – vi läser bara XML)."""
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = _VARIANT_RE[kind].sub(_MAIN_TYPES[kind], data.decode("utf-8")).encode("utf-8")
            dst.writestr(item, data)
    buf.seek(0)
    buf.name = Path(path).stem + "." + kind
    return buf


def extract_office(path: Path, *, images_dir: Path | None = None, max_rows: int = 2000,
                   password: str | None = None, figures_dir: Path | None = None) -> tuple[str, dict]:
    ext = Path(path).suffix.lower()
    stem = Path(path).stem
    fd = figures_dir
    if ext == ".docx":
        return extract_docx(path, images_dir, password, figures_dir=fd)
    if ext in (".docm", ".dotx", ".dotm"):
        return extract_docx(normalize_ooxml(path, "docx"), images_dir, stem=stem, figures_dir=fd)
    if ext in (".doc", ".rtf", ".odt", ".wpd", ".dot", ".wps"):
        md, meta = extract_docx(convert_legacy(path, "docx"), images_dir, stem=stem, figures_dir=fd)
        meta["converted_from"] = ext
        return md, meta
    if ext in (".xlsx", ".xlsm", ".xltx", ".xls", ".ods"):
        return extract_sheet(path, max_rows, password)
    if ext == ".pptx":
        return extract_pptx(path, password, figures_dir=fd)
    if ext in (".pptm", ".ppsx", ".potx", ".ppsm"):
        return extract_pptx(normalize_ooxml(path, "pptx"), figures_dir=fd, stem=stem)
    if ext in (".ppt", ".odp", ".pps"):
        md, meta = extract_pptx(convert_legacy(path, "pptx"), figures_dir=fd, stem=stem)
        meta["converted_from"] = ext
        return md, meta
    raise ValueError(f"Inte ett Office-format: {ext}")
